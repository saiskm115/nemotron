import io
from typing import List, Dict
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from ..models.session import Session

def export_docx(session: Session, include_translation: bool = True) -> bytes:
    doc = docx.Document()

    # Document Title
    title = doc.add_heading(session.title, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Metadata Paragraph
    meta = doc.add_paragraph()
    meta.add_run(f"Session ID: ").bold = True
    meta.add_run(f"{session.id} | ")
    meta.add_run(f"Duration: ").bold = True
    mins = int(session.duration // 60)
    secs = int(session.duration % 60)
    meta.add_run(f"{mins:02d}:{secs:02d} | ")
    meta.add_run(f"Speakers: ").bold = True
    meta.add_run(f"{len(session.speakers)} | ")
    meta.add_run(f"Languages: ").bold = True
    meta.add_run(", ".join(session.source_languages))

    doc.add_heading("Speakers", level=2)
    spk_table = doc.add_table(rows=1, cols=3)
    hdr_cells = spk_table.rows[0].cells
    hdr_cells[0].text = "Speaker"
    hdr_cells[1].text = "Model Label"
    hdr_cells[2].text = "Speaking Time"
    for s in session.speakers:
        row = spk_table.add_row().cells
        row[0].text = s.display_name
        row[1].text = s.model_label
        s_dur = f"{int(s.total_speaking_time // 60):02d}:{int(s.total_speaking_time % 60):02d}"
        row[2].text = s_dur

    doc.add_paragraph()
    doc.add_heading("Transcript", level=2)

    spk_map: Dict[str, str] = {s.id: s.display_name for s in session.speakers}

    for turn in session.turns:
        spk_name = spk_map.get(turn.speaker_id, turn.speaker_id)
        start_min = int(turn.start // 60)
        start_sec = int(turn.start % 60)
        time_tag = f"[{start_min:02d}:{start_sec:02d}]"

        p = doc.add_paragraph()
        run_spk = p.add_run(f"{spk_name} {time_tag}: ")
        run_spk.bold = True
        run_spk.font.color.rgb = RGBColor(14, 116, 144) # Deep teal

        if turn.overlap:
            run_ov = p.add_run(" [Overlapping Speech] ")
            run_ov.font.color.rgb = RGBColor(220, 38, 38)
            run_ov.italic = True

        p.add_run(f"{turn.text}\n")

        if include_translation and turn.translated_text:
            p_trans = doc.add_paragraph()
            p_trans.paragraph_format.left_indent = Inches(0.25)
            run_tr_lbl = p_trans.add_run("Translation: ")
            run_tr_lbl.italic = True
            run_tr_lbl.font.color.rgb = RGBColor(100, 116, 139)
            run_tr = p_trans.add_run(turn.translated_text)
            run_tr.italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
